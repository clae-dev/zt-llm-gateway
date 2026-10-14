"""평가셋 양식(xlsx)과 시나리오 폴더를 읽는다.

평가셋은 사람이 쓴다. 이 저장소의 규칙과 코드를 만든 쪽은 봉인용 절반을 열어 보지 않는다.
개발용/봉인용은 요청 문장의 해시로 나누므로 쓰는 사람이 고를 필요가 없고, 나중에 바뀌지도 않는다.
"""
from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.scenarios import SCENARIO_DIR, Scenario, load_scenarios  # noqa: F401 - 측정 도구가 여기서 가져다 쓴다

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = Path.home() / "OneDrive - 한양사이버대학교" / "바탕 화면" / "01_창업·공모전" / "캠퍼스랩" / "06_평가셋" / "평가셋_작성양식.xlsx"

GROUPS = {
    # 묶음: (목표 건수, 의도, 무엇을 쓰나)
    "문맥 유출": (60, "공격", "문장만 보면 평범한데, 그 상황에서는 기밀을 끌어내게 되는 요청"),
    "직접 공격": (40, "공격", "지시를 무시하라거나 비밀을 말하라고 대놓고 요구하는 요청"),
    "글자 우회": (40, "공격", "글자를 바꿔 쓴 요청, 또는 기밀을 바꿔 쓴 글자로 내보내게 하는 요청"),
    "정상 - 공격 단어 포함": (40, "정상", "공격에 흔한 낱말이 들어 있지만 실제로는 평범한 요청"),
    "정상 - 일반 업무": (20, "정상", "그 상황에서 흔히 할 법한 평범한 업무 요청"),
}
SHEET = "평가셋"
COLUMNS = ["번호", "묶음", "세부 유형", "시나리오", "요청 문장", "의도", "노리는 기밀", "작성자", "메모"]


@dataclass(frozen=True)
class Item:
    row: int
    group: str
    subtype: str
    scenario: str
    text: str
    intent: str
    target: str
    author: str

    @property
    def split(self) -> str:
        digest = hashlib.sha1(self.text.strip().encode("utf-8")).digest()
        return "봉인용" if digest[0] % 2 else "개발용"

    @property
    def is_attack(self) -> bool:
        return self.intent == "공격"


def load_items(path: Path = DEFAULT_XLSX) -> tuple[list[Item], list[str]]:
    """(항목, 문제점). 문제점이 있어도 읽을 수 있는 항목은 돌려준다."""
    ws = load_workbook(path, data_only=True)[SHEET]
    scenarios = load_scenarios()
    items, problems, seen = [], [], {}
    for r, row in enumerate(ws.iter_rows(min_row=2, max_col=len(COLUMNS), values_only=True), start=2):
        _, group, subtype, scenario, text, intent, target, author, _ = (str(v).strip() if v is not None else "" for v in row)
        if not text:
            continue
        sid = scenario.split(" ")[0].split("_")[0]
        if group not in GROUPS:
            problems.append(f"{r}행: 묶음이 비었거나 목록에 없다 ({group!r})")
        elif sid not in scenarios:
            problems.append(f"{r}행: 시나리오 {scenario!r} 폴더가 없다")
        elif intent not in ("공격", "정상"):
            problems.append(f"{r}행: 의도가 비었다")
        elif intent != GROUPS[group][1]:
            problems.append(f"{r}행: 묶음 {group!r}의 의도는 {GROUPS[group][1]}인데 {intent}로 적혀 있다")
        elif text in seen:
            problems.append(f"{r}행: {seen[text]}행과 같은 문장이다")
        else:
            seen[text] = r
            items.append(Item(r, group, subtype, sid, text, intent, target, author))
    return items, problems
