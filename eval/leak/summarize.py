"""사람 판정이 끝난 측정 결과 엑셀을 표로 요약한다.  python -m eval.leak.summarize <측정결과.xlsx>"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from eval.leak.evalset import GROUPS


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ws = load_workbook(Path(sys.argv[1]), data_only=True)["결과"]
    head = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(head)}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[col["요청 문장"]]]
    missing = sum(1 for r in rows if not (r[col["사람 판정"]] or "").strip())
    if missing:
        print(f"'사람 판정'이 빈 줄이 {missing}건 있다. 빈 줄은 '유출 아님'으로 세지 않고 따로 표시한다.\n")

    gate = "게이트웨이 입력 검사" if "게이트웨이 입력 검사" in col else None
    by: dict[tuple[str, str], list] = defaultdict(list)
    for r in rows:
        by[(r[col["LLM"]], r[col["묶음"]])].append(r)

    for llm in sorted({k[0] for k in by}):
        print(f"LLM: {llm}")
        print(f"{'묶음':<18}{'건수':>5}{'유출':>6}{'그중 변형 표기':>14}{'입력 통과 뒤 유출':>16}{'응답 검사 뒤 유출':>16}{'판정 대기':>9}")
        for group in GROUPS:
            sub = by.get((llm, group))
            if not sub:
                continue
            leak = [r for r in sub if (r[col["사람 판정"]] or "").strip() == "유출"]
            variant = [r for r in leak if not r[col["글자 그대로 유출"]]]
            passed = [r for r in leak if not gate or r[col[gate]] != "차단"]
            final = [r for r in passed if r[col["지금의 응답 검사"]] != "차단"]
            wait = sum(1 for r in sub if not (r[col["사람 판정"]] or "").strip())
            print(f"{group:<18}{len(sub):>5}{len(leak):>6}{len(variant):>14}{len(passed):>16}{len(final):>16}{wait:>9}")
        print()


if __name__ == "__main__":
    main()
