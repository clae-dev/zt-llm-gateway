"""측정 결과 엑셀들을 모아 숫자만 집계한다. 문장과 답은 읽어서 찍지 않는다.

python datasets/external/aihub_102/summarize_big.py
  150건(첫 측정)과 big/c1~c5(1,000건)를 모델별·나눔별·의도별로 센다.
"""
import collections
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
sys.stdout.reconfigure(encoding="utf-8")


def rows_of(path: Path) -> list[dict]:
    ws = load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    return [dict(zip(head, r)) for r in it]


def report(title: str, files: list[Path]) -> None:
    if not files:   # 측정 꾸러미에는 기존 결과 파일이 없다
        return
    stat = collections.defaultdict(lambda: collections.Counter())
    for f in files:
        split = "봉인용" if "봉인용" in f.name else "개발용"
        for r in rows_of(f):
            llm = r["LLM"]
            top = str(r["세부 유형"]).split("_")[0]
            leak = bool(r["글자 그대로 유출"])
            blocked = r["지금의 응답 검사"] == "차단"
            in_block = r.get("게이트웨이 입력 검사") == "차단"
            for key in ((llm, "전체"), (llm, split), (llm, "의도 " + top)):
                c = stat[key]
                c["건수"] += 1
                c["유출"] += leak
                c["유출 뒤 응답 검사 통과"] += leak and not blocked
                c["정상 답인데 응답 차단"] += blocked and not leak
                c["입력 오차단"] += in_block
            if leak:
                for kind in str(r["글자 그대로 유출"]).split(", "):
                    stat[(llm, "나온 것 " + kind)]["건수"] += 1
    print(f"\n== {title} (파일 {len(files)}개) ==")
    for (llm, part), c in sorted(stat.items()):
        if part.startswith("나온 것"):
            print(f"  {llm:<16}{part:<22}{c['건수']:>5}")
        else:
            print(f"  {llm:<16}{part:<22}건수 {c['건수']:>5}  유출 {c['유출']:>4}  남은 유출 {c['유출 뒤 응답 검사 통과']:>3}"
                  f"  응답 오차단 {c['정상 답인데 응답 차단']:>3}  입력 오차단 {c['입력 오차단']:>3}")


report("첫 150건 · EXAONE", sorted(HERE.glob("측정결과_*.xlsx")))
report("첫 150건 · gemma3", sorted((HERE / "gemma3").glob("측정결과_*.xlsx")))
report("추가 1,000건 · 두 모델", sorted((HERE / "big").glob("c*/측정결과_*.xlsx")))
report("같은 1,150건 · EXAONE 2.4B", sorted((HERE / "small").glob("*/측정결과_*.xlsx")))
for model in sorted(p for p in (HERE / "npu").glob("*") if p.is_dir()):   # run_models.py 결과(NPU 서버 모델별)
    report(f"같은 1,150건 · NPU {model.name}", sorted(model.glob("*/측정결과_*.xlsx")))
print("\n진행 기록:")
status = HERE / "big" / "status.txt"
if status.exists():
    print("".join(status.read_text(encoding="utf-8").splitlines(keepends=True)[-6:]))
