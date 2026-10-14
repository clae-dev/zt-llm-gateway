"""AI Hub '소상공인 고객 주문 질의-응답 텍스트'에서 S2(쇼핑몰 상담)용 고객 질문 1,000건을 더 뽑는다.

sample_s2.py와 같은 규칙이다. 사람이 고르지 않고, 문장은 화면에 찍지 않는다.
  원본   : 디지털가전_validation.csv (발화자 c, QA여부 q)
  의도   : 배송 400, 교환|반품|환불 400, 결제 100, AS 100  (150건 때와 같은 비율 4:4:1:1)
  제외   : 상품명 칸이 채워진 줄, 10자 미만·80자 초과, '#' 가림 표시가 든 줄, 같은 문장,
           그리고 먼저 뽑은 150건(평가셋_AIHub_S2.xlsx)에 든 문장
  시드   : 20261015
  나눔   : 중간에 멈춰도 결과가 남도록 200건씩 5묶음(big/c1 ~ c5)으로 저장한다
출처: 한국지능정보사회진흥원 AI Hub. 원문을 팀 밖에 공유하거나 저장소에 올리지 않는다.
"""
import csv
import random
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
QUOTA = {"배송": 400, "교환|반품|환불": 400, "결제": 100, "AS": 100}
COLUMNS = ["번호", "묶음", "세부 유형", "시나리오", "요청 문장", "의도", "노리는 기밀", "작성자", "메모"]
CHUNKS = 5

used = {str(r[4]).strip() for r in load_workbook(HERE / "평가셋_AIHub_S2.xlsx")["평가셋"].iter_rows(min_row=2, values_only=True) if r[4]}
rows = list(csv.DictReader(open(HERE / "디지털가전_validation.csv", encoding="utf-8")))
pool: dict[str, list[dict]] = {k: [] for k in QUOTA}
seen = set(used)
for r in rows:
    text = r["발화문"].strip()
    top = r["인텐트"].split("_")[0]
    if (r["발화자"] == "c" and r["QA여부"] == "q" and top in QUOTA and not r["상품명"].strip()
            and 10 <= len(text) <= 80 and "#" not in text and text not in seen):
        seen.add(text)
        pool[top].append(r)

rng = random.Random(20261015)
picked = []
for top, k in QUOTA.items():
    picked += rng.sample(pool[top], k)
rng.shuffle(picked)  # 묶음마다 의도가 고르게 섞이도록

size = len(picked) // CHUNKS
for c in range(CHUNKS):
    wb = Workbook()
    ws = wb.active
    ws.title = "평가셋"
    ws.append(COLUMNS)
    for n, r in enumerate(picked[c * size:(c + 1) * size], 1):
        ws.append([n, "정상 - 일반 업무", r["인텐트"], "S2", r["발화문"].strip(), "정상", "",
                   "AI Hub 소상공인 고객 주문 질의-응답 (실제 고객 발화)", f"IDX {r['IDX']}"])
    folder = HERE / "big" / f"c{c + 1}"
    folder.mkdir(parents=True, exist_ok=True)
    wb.save(folder / "평가셋_AIHub_S2.xlsx")
print({k: len(v) for k, v in pool.items()}, "->", len(picked), "건,", CHUNKS, "묶음, 먼저 쓴", len(used), "건 제외")
